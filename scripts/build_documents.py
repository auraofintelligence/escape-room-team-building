"""Create the portable business plan from the canonical website manuscript.

Run with the primary artifact Python runtime. The PDF is made by the documents
skill's render_docx.py so Word and PDF share the same layout and manuscript.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import zipfile
from io import BytesIO
from pathlib import Path
from urllib.parse import quote, urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfbase.pdfmetrics import stringWidth


ROOT = Path(__file__).resolve().parents[1]
BASE = "https://auraofintelligence.github.io/escape-room-team-building/"
NAVY = "203B4C"
SAGE = "EDF2EE"
LINK = "275E67"
GRAY = "526568"
BLACK = "000000"
FONT = "Arial"
CONTENT_WIDTH = 7.05
TABLE_FONT = 10.2
_bookmark_id = 1


def element(name, attrs=None):
    node = OxmlElement(name)
    for k, v in (attrs or {}).items():
        node.set(qn(k), str(v))
    return node


def bookmark(paragraph, name):
    global _bookmark_id
    paragraph._p.insert(0, element("w:bookmarkStart", {"w:id": _bookmark_id, "w:name": name}))
    paragraph._p.append(element("w:bookmarkEnd", {"w:id": _bookmark_id}))
    _bookmark_id += 1


def font_style(style, size, *, bold=False, color=BLACK):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = element("w:rFonts")
        rpr.insert(0, fonts)
    for name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(name), FONT)
    for node in rpr.findall(qn("w:color")):
        for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
            node.attrib.pop(qn(attr), None)


def add_style(doc, name, size, **kwargs):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles["Normal"]
    font_style(s, size, **kwargs)
    return s


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(.725)
    sec.right_margin = Inches(.725)
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.header_distance = Inches(.22)
    sec.footer_distance = Inches(.25)
    normal = doc.styles["Normal"]
    font_style(normal, 11)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 21, 0, 5), ("Subtitle", 11, 0, 5), ("Heading 1", 21, 0, 12), ("Heading 2", 14, 16, 7), ("Heading 3", 11.8, 12, 6)]:
        s = doc.styles[name]
        font_style(s, size, bold=name != "Subtitle")
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.line_spacing = 1.06
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.keep_together = True
        ppr = s.element.get_or_add_pPr()
        for tag in ("w:pBdr", "w:shd"):
            for n in ppr.findall(qn(tag)):
                ppr.remove(n)
    for name, size, color, bold in [("Lead", 13.4, BLACK, False), ("Metadata", 9.5, GRAY, False), ("SourceNote", 9.2, GRAY, False), ("TableText", TABLE_FONT, BLACK, False), ("TableHead", TABLE_FONT, "FFFFFF", True), ("Contents", 11.1, BLACK, False), ("SourceTitle", 11, BLACK, True), ("RegisterText", 10.2, BLACK, False)]:
        s = add_style(doc, name, size, color=color, bold=bold)
        s.paragraph_format.line_spacing = 1.10
        s.paragraph_format.space_after = Pt(6)
        if name in ("SourceTitle", "TableHead"):
            s.paragraph_format.keep_with_next = True
        if name == "SourceTitle":
            s.paragraph_format.space_before = Pt(10)
        if name in ("TableText", "TableHead"):
            s.paragraph_format.space_after = Pt(0)
            s.paragraph_format.line_spacing = 1.06
        if name == "Contents":
            s.paragraph_format.space_after = Pt(11)
            s.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    for name in ("List Bullet", "List Number"):
        font_style(doc.styles[name], 11)
        doc.styles[name].paragraph_format.line_spacing = 1.10
        doc.styles[name].paragraph_format.space_after = Pt(5)
    sec.footer.paragraphs[0].style = doc.styles["Metadata"]
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run("Escape Room Team Building  |  Luke Nathan Hayes\t")
    field(footer, "PAGE", "1")
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = element("w:updateFields", {"w:val": "true"})
        settings.append(update)
    lang = doc.styles["Normal"].element.get_or_add_rPr()
    lang.append(element("w:lang", {"w:val": "en-AU"}))


def field(paragraph, instruction, display):
    begin = paragraph.add_run()._r
    begin.append(element("w:fldChar", {"w:fldCharType": "begin"}))
    ins = paragraph.add_run()._r
    t = element("w:instrText", {"xml:space": "preserve"})
    t.text = " " + instruction + " "
    ins.append(t)
    paragraph.add_run()._r.append(element("w:fldChar", {"w:fldCharType": "separate"}))
    paragraph.add_run(str(display))
    paragraph.add_run()._r.append(element("w:fldChar", {"w:fldCharType": "end"}))


def hyperlink(paragraph, label, target, *, internal=False, color=LINK, bold=False, size=None):
    h = element("w:hyperlink", {"w:anchor": target} if internal else {"r:id": paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)})
    r = element("w:r")
    rp = element("w:rPr")
    rp.append(element("w:color", {"w:val": color}))
    rp.append(element("w:u", {"w:val": "single"}))
    if bold:
        rp.append(element("w:b"))
    if size:
        rp.append(element("w:sz", {"w:val": round(size * 2)}))
    r.append(rp)
    t = element("w:t", {"xml:space": "preserve"})
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


TOKEN = re.compile(r"(\[[^\]]+\]\([^\s)]+\)|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\b[DPMR]\d{2,3}\b|https?://[^\s<>]+)")


def inline(paragraph, text, refs, *, color=None):
    for token in TOKEN.split(str(text)):
        if not token:
            continue
        if token.startswith("[") and "](" in token:
            label, target = token[1:].split("](", 1)
            target = target[:-1]
            if target.startswith("#"):
                hyperlink(paragraph, label, target[1:], internal=True, color=color or LINK)
            else:
                if not urlsplit(target).scheme:
                    target = BASE + target
                hyperlink(paragraph, label, target, color=color or LINK)
        elif token in refs:
            hyperlink(paragraph, token, "source_" + token, internal=True, color=color or LINK)
        elif token.startswith("https://") or token.startswith("http://"):
            end = ""
            while token and token[-1] in ".,;":
                end = token[-1] + end
                token = token[:-1]
            hyperlink(paragraph, token, token, color=color or LINK)
            paragraph.add_run(end)
        else:
            r = paragraph.add_run()
            if token.startswith("**") and token.endswith("**"):
                r.text, r.bold = token[2:-2], True
            elif token.startswith("*") and token.endswith("*"):
                r.text, r.italic = token[1:-1], True
            elif token.startswith("`") and token.endswith("`"):
                r.text = token[1:-1]
                r.font.name = "Courier New"
                r.font.size = Pt(10)
            else:
                r.text = token
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def table_cells(text):
    return [x.strip().replace("\\|", "|") for x in re.split(r"(?<!\\)\|", text.strip().strip("|"))]


def widths_for(headers, rows):
    n = len(headers)
    samples = [[re.sub(r"\*|`|\[[^\]]+\]\([^)]*\)", "", str(r[i])) for r in [headers] + rows if i < len(r)] for i in range(n)]
    # Paragraph density controls column width; numeric columns remain compact.
    weights = []
    numeric = []
    for i, col in enumerate(samples):
        vals = col[1:]
        num = bool(vals) and sum(bool(re.fullmatch(r"[\d\s,$.%+()×÷\-–/]+", s)) for s in vals) / len(vals) > .7
        numeric.append(num)
        if num:
            weights.append(max(8, len(headers[i]) * .65, min(18, max(map(len, col)) * .72)))
        else:
            sorted_lens = sorted(map(len, col))
            typical = sorted_lens[min(len(sorted_lens)-1, int(len(sorted_lens)*.75))]
            weights.append(max(12, min(50, math.sqrt(typical) * 4.3)))
    if n == 2:
        weights = [min(weights[0], weights[1]*.70), weights[1]]
    total = sum(weights)
    widths = [CONTENT_WIDTH * w / total for w in weights]
    minimum = .66 if n > 5 else .82
    minima = [max(minimum, .075 * max(map(len, re.findall(r"[A-Za-z]+", h) or [""])) + .19) for h in headers]
    for i, is_number in enumerate(numeric):
        if is_number:
            numeric_values = [v for v in samples[i][1:] if re.fullmatch(r"[\d\s,$.%+()×÷\-–/]+", v)]
            numeric_width = max(stringWidth(v, "Helvetica", TABLE_FONT) for v in numeric_values)
            minima[i] = max(minima[i], (numeric_width + 12) / 72)
    if sum(minima) > CONTENT_WIDTH:
        raise ValueError("Table needs a wider print layout: " + str(headers))
    for _ in range(3):
        extra = sum(max(0, m-w) for m,w in zip(minima,widths))
        if not extra:
            break
        surplus = sum(max(0,w-m) for m,w in zip(minima,widths))
        widths = [m if w < m else w-extra*(w-m)/surplus for m,w in zip(minima,widths)]
    return widths, numeric


def add_table(doc, headers, rows, refs):
    rows = [list(r) + [""] * (len(headers)-len(r)) for r in rows]
    if any(len(r) != len(headers) for r in rows):
        raise ValueError("Table row does not match headers: " + str(headers))
    widths, numeric = widths_for(headers, rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    pr = table._tbl.tblPr
    tw = pr.find(qn("w:tblW"))
    tw.set(qn("w:w"), str(round(CONTENT_WIDTH * 1440)))
    tw.set(qn("w:type"), "dxa")
    pr.append(element("w:tblLayout", {"w:type": "fixed"}))
    borders = element("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element("w:"+edge, {"w:val": "single", "w:sz": "5", "w:color": "D9D9D9"}))
    pr.append(borders)
    margins = element("w:tblCellMar")
    for edge, val in (("top", 85), ("left", 95), ("bottom", 85), ("right", 95)):
        margins.append(element("w:"+edge, {"w:w": val, "w:type": "dxa"}))
    pr.append(margins)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)
    for row_index, data in enumerate([headers] + rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(element("w:cantSplit"))
        if row_index == 0:
            trpr.append(element("w:tblHeader", {"w:val": "true"}))
        for i, value in enumerate(data):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if row_index == 0 or numeric[i] else WD_CELL_VERTICAL_ALIGNMENT.TOP
            fill = NAVY if row_index == 0 else SAGE if row_index % 2 else "FFFFFF"
            cell._tc.get_or_add_tcPr().append(element("w:shd", {"w:fill": fill}))
            p = cell.paragraphs[0]
            p.style = doc.styles["TableHead" if row_index == 0 else "TableText"]
            p.paragraph_format.keep_with_next = row_index == 0
            p.paragraph_format.keep_together = True
            if numeric[i]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value, refs, color="FFFFFF" if row_index == 0 else None)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(3)
    return table


def markdown(doc, text, refs, *, heading_offset=0, style=None):
    lines = str(text).strip().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if "|" in line and i+1 < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[i+1]) and "-" in lines[i+1]:
            headers = table_cells(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(table_cells(lines[i]))
                i += 1
            add_table(doc, headers, rows, refs)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = min(3, len(heading[1]) + heading_offset)
            p = doc.add_paragraph(style="Heading " + str(max(1, level)))
            inline(p, heading[2], refs)
            i += 1
            continue
        li = re.match(r"^([-*]|\d+[.)])\s+(.+)$", line)
        if li:
            ordered = li[1][0].isdigit()
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            # Preserve explicit ordered-list numbering rather than continuing
            # an unrelated Word list elsewhere in the document.
            if ordered:
                p.style = doc.styles["Normal"]
                p.paragraph_format.left_indent = Inches(.20)
                p.paragraph_format.first_line_indent = Inches(-.20)
                inline(p, li[1] + " " + li[2], refs)
            else:
                inline(p, li[2], refs)
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style=style or "Normal")
            inline(p, line[2:], refs)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(?:[-*] |\d+[.)] |#{1,6} |> )", lines[i]):
            if "|" in lines[i] and i+1 < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[i+1]):
                break
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph(style=style or "Normal")
        inline(p, " ".join(para), refs)


def reference_note(doc, ids, refs):
    ids = list(dict.fromkeys(ids))
    if ids:
        p = doc.add_paragraph(style="SourceNote")
        p.add_run("Sources: ")
        inline(p, ", ".join(ids), refs)


def chapter(doc, c, refs, first=False):
    p = doc.add_paragraph(style="Heading 1")
    if not first and c["slug"] not in {"renewal", "partnerships", "roadmap"}:
        p.paragraph_format.page_break_before = True
    elif not first:
        p.paragraph_format.space_before = Pt(18)
    title = "EXECUTIVE SUMMARY" if first else c["title"]
    inline(p, title, refs)
    bookmark(p, "chapter_" + re.sub(r"\W", "_", c["slug"]))
    p = doc.add_paragraph(style="Lead")
    inline(p, c["lead"], refs)
    if first:
        for r in p.runs:
            r.bold = True
    for s in c["sections"]:
        h = doc.add_paragraph(s["title"], "Heading 2")
        raw_name = c["slug"] + "_" + s.get("id", s["title"])
        name = re.sub(r"\W", "_", raw_name)[:29] + "_" + hashlib.sha256(raw_name.encode()).hexdigest()[:6]
        bookmark(h, "s_" + name)
        if s.get('basis'):
            doc.add_paragraph(s['basis'], style='SourceNote').paragraph_format.keep_with_next = True
        markdown(doc, s["body"], refs)
        reference_note(doc, s.get("refs", []), refs)
    reference_note(doc, c.get("refs", []), refs)


def contents(doc, plan, pagination, follows_art=False):
    p = doc.add_paragraph("CONTENTS", "Heading 1")
    p.paragraph_format.page_break_before = not follows_art
    if follows_art:
        p.paragraph_format.space_before = Pt(14)
    bookmark(p, "contents")
    entries = [(c["title"], "chapter_"+re.sub(r"\W", "_", c["slug"])) for c in plan["chapters"]]
    entries += [("Sources and project inventory", "source_register"), ("Rights and provenance", "rights"), ("Strange But True Public Source Licence", "licence")]
    for label, target in entries:
        p = doc.add_paragraph(style="Contents")
        hyperlink(p, label, target, internal=True)
    p = doc.add_paragraph(style="SourceNote")
    hyperlink(p, "Online business plan", BASE)
    p.add_run("  |  ")
    hyperlink(p, "Source files and full reference trail", BASE+"sources.html")


def opening_art(doc):
    path = ROOT/"content/artwork.json"
    if not path.is_file():
        return False
    entry = next((a for a in json.loads(path.read_text(encoding="utf-8")) if a.get("slug") == "index"), None)
    if not entry or not (ROOT/entry["path"]).is_file():
        return False
    from PIL import Image
    png = BytesIO()
    with Image.open(ROOT/entry["path"]) as art:
        art.convert("RGB").save(png, format="PNG")
    png.seek(0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = p.add_run().add_picture(png, width=Inches(5.2))
    pic._inline.docPr.set("descr", entry.get("alt", "Concept artwork for the proposed venue"))
    caption = doc.add_paragraph(entry.get("caption", "Original AI-generated concept artwork for this business plan."), style="SourceNote")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    return True


def source_register(doc, sources, refs):
    p = doc.add_paragraph("Sources and project inventory", "Heading 1")
    p.paragraph_format.page_break_before = True
    bookmark(p, "source_register")
    p = doc.add_paragraph()
    inline(p, "The original documents and project websites from Installing the Missing Middle inform the business design. Market and cost references distinguish source observations, worked assumptions and verification status. Source IDs below match the website and canonical source register.", refs)
    p = doc.add_paragraph(style="SourceNote")
    hyperlink(p, "Complete source register", BASE+"content/sources.json")
    p.add_run("  |  ")
    hyperlink(p, "Full inherited embedded-link inventory", BASE+"content/embedded-links.json")
    for key, heading in [("documents", "Original documents"), ("links", "Connected projects and websites"), ("research", "Market and operating sources")]:
        if not sources.get(key):
            continue
        doc.add_paragraph(heading, "Heading 2")
        for entry in sources.get(key, []):
            p = doc.add_paragraph(style="SourceTitle")
            p.add_run(entry["id"] + "  " + entry["title"])
            bookmark(p, "source_"+entry["id"])
            note = entry.get("application", entry.get("purpose", entry.get("note", entry.get("support", ""))))
            if note:
                markdown(doc, note, refs, style="RegisterText")
            if entry.get("checked"):
                p = doc.add_paragraph("Recorded observation date: " + str(entry["checked"]), style="SourceNote")
                p.paragraph_format.keep_with_next = True
            if entry.get('verification_status'):
                doc.add_paragraph(entry['verification_status'].replace('_',' ')+': '+entry.get('verification_note',''), style='SourceNote')
                doc.add_paragraph('Supplier quotation date: none recorded.', style='SourceNote')
            p = doc.add_paragraph(style="SourceNote")
            if key == "documents":
                path = entry.get("path", "")
                if path and (ROOT/path).is_file():
                    hyperlink(p, "Original " + Path(path).suffix[1:].upper(), BASE+quote(path, safe="/"))
                    if (ROOT/"documents/readers"/f'{entry["id"]}.html').is_file():
                        p.add_run("  |  ")
                        hyperlink(p, "Online reader", BASE+f'documents/readers/{entry["id"]}.html')
                    if entry.get("source_url"):
                        p.add_run("  |  ")
                        hyperlink(p, "Archived source", entry["source_url"])
                else:
                    p.add_run("Exact original unavailable in the inherited archive. See the source entry for the available edition.")
            elif entry.get("url"):
                url = entry["url"]
                # Print the original URL as well as making it clickable.
                hyperlink(p, url, url)
            else:
                hyperlink(p, "Source entry", BASE+"sources.html#"+entry["id"])


def validate_inputs(plan, sources):
    if not plan.get("chapters"):
        raise ValueError("No canonical chapters found")
    text = json.dumps(plan, ensure_ascii=False)
    if "\u2014" in text:
        raise ValueError("Em dash found in the canonical manuscript")
    all_ids = [s["id"] for k in ("documents", "links", "research") for s in sources.get(k, [])]
    if len(all_ids) != len(set(all_ids)):
        raise ValueError("Duplicate source IDs")
    mentioned = set()
    for c in plan["chapters"]:
        mentioned.update(c.get("refs", []))
        for s in c["sections"]:
            mentioned.update(s.get("refs", []))
            mentioned.update(re.findall(r"\b[DPMR]\d{2,3}\b", s["body"]))
    missing = mentioned-set(all_ids)
    if missing:
        raise ValueError("Undefined source IDs: " + ", ".join(sorted(missing)))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--plan", type=Path, default=ROOT/"content/plan.json")
    parser.add_argument("--sources", type=Path, default=ROOT/"content/sources.json")
    parser.add_argument("--output", type=Path, default=ROOT/"downloads/Escape-Room-Network-Business-Plan.docx")
    parser.add_argument("--pagination-json", type=Path)
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text(encoding="utf-8"))
    sources = json.loads(args.sources.read_text(encoding="utf-8"))
    validate_inputs(plan, sources)
    refs = {e["id"]:e for key in ("documents", "links", "research") for e in sources.get(key, [])}
    pagination = json.loads(args.pagination_json.read_text(encoding="utf-8")) if args.pagination_json else {}
    doc = Document()
    setup(doc)
    doc.core_properties.title = plan.get("title", "Escape Room Team Building Business Plan")
    doc.core_properties.author = plan.get("author", "Luke Nathan Hayes")
    doc.core_properties.subject = "Australian escape-room network business plan"
    doc.core_properties.keywords = "escape rooms, team building, tourism, whole-of-life learning, Australia"
    doc.core_properties.comments = ""
    title = doc.add_paragraph(plan.get("title", "Escape Room Team Building Business Plan"), "Title")
    title.paragraph_format.keep_with_next = True
    if plan.get("subtitle"):
        doc.add_paragraph(plan["subtitle"], "Subtitle")
    p = doc.add_paragraph(style="Metadata")
    p.add_run(plan.get("author", "Luke Nathan Hayes") + "  |  " + plan.get("date", "29 October 2026") + "  |  " + plan.get("edition", "Business plan 1.0"))
    p.paragraph_format.keep_with_next = True
    chapter(doc, plan["chapters"][0], refs, first=True)
    has_opening_art = opening_art(doc)
    contents(doc, plan, pagination, has_opening_art)
    for c in plan["chapters"][1:]:
        chapter(doc, c, refs)
    source_register(doc, sources, refs)
    for file, title, mark in [("RIGHTS.md", "Rights and provenance", "rights"), ("LICENCE.md", "Strange But True Public Source Licence", "licence")]:
        p = doc.add_paragraph(title, "Heading 1")
        p.paragraph_format.space_before = Pt(18)
        bookmark(p, mark)
        text = (ROOT/file).read_text(encoding="utf-8")
        # The heading already appears above; preserve every substantive term.
        text = re.sub(r"\A# [^\n]+\n+", "", text)
        first_body_paragraph = len(doc.paragraphs)
        markdown(doc, text, refs)
        if file == "LICENCE.md":
            # Keep the full licence together without a short final-page spillover.
            # Body type remains 11 pt; only the appendix's paragraph gaps tighten.
            for licence_paragraph in doc.paragraphs[first_body_paragraph:]:
                if licence_paragraph.style.name.startswith("Heading"):
                    licence_paragraph.paragraph_format.space_before = Pt(12)
                else:
                    licence_paragraph.paragraph_format.space_after = Pt(5)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = args.output.with_name("." + args.output.name + ".building")
    doc.save(temporary_output)
    with zipfile.ZipFile(temporary_output) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX archive failed integrity verification")
    temporary_output.replace(args.output)
    report = {
        "docx": str(args.output.relative_to(ROOT)) if args.output.is_relative_to(ROOT) else str(args.output),
        "canonical_plan_sha256": hashlib.sha256(args.plan.read_bytes()).hexdigest(),
        "source_register_sha256": hashlib.sha256(args.sources.read_bytes()).hexdigest(),
        "chapters": len(plan["chapters"]),
        "sections": sum(len(c["sections"]) for c in plan["chapters"]),
        "source_entries": len(refs),
        "table_count": len(doc.tables),
        "manuscript_words": len(re.findall(r"\S+", " ".join(c["title"]+" "+c["lead"]+" "+" ".join(s["title"]+" "+s["body"] for s in c["sections"]) for c in plan["chapters"]))),
        "licence_included": True,
        "all_chapter_bodies_included": True,
        "research_status": "ROUND ONE exploratory comparisons for founder review"
    }
    (ROOT/"content/document-build-report.json").write_text(json.dumps(report, indent=2)+"\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
